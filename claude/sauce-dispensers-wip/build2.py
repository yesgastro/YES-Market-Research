"""Append a product family (offers + benchmarks + review + compare + pricing + change log) to the workbook.
usage: build2.py <src.xlsx> <offers.csv> <spec.json> <out.xlsx> [changelog_extra.csv]"""
import sys, re, csv, copy, json, datetime
from openpyxl import load_workbook
from collections import Counter, defaultdict

SRC, CSV, SPEC, OUT = sys.argv[1:5]
EXTRA = sys.argv[5] if len(sys.argv) > 5 else None
spec = json.load(open(SPEC, encoding='utf-8'))
TODAY = '2 Oct 2026'; DATE_ISO = '2026-10-02'; DT = datetime.datetime(2026, 10, 2)
OFF_MAX, REV_MAX = 5000, 1500
MARKETS = ['Hungary', 'Romania', 'Austria', 'Germany', 'Czech', 'Slovakia']
CUR = {'Hungary': 'HUF', 'Romania': 'RON', 'Austria': 'EUR', 'Germany': 'EUR', 'Czech': 'CZK', 'Slovakia': 'EUR'}
NEW = [(v['product'], v['code'], v['spec'], v['req']) for v in spec['variants']]
NEW_NAMES = [n[0] for n in NEW]
FAMILY = spec['family']; FAMILY_SHORT = spec['family_short']

rows = list(csv.DictReader(open(CSV, encoding='utf-8')))
for r in rows:
    assert r['product'] in NEW_NAMES, ('unknown product', r['product']); assert r['market'] in MARKETS, r['market']
    r['price_shown'] = float(r['price_shown']); r['pack_size'] = int(float(r['pack_size'] or 1)); r['vat_rate'] = float(r['vat_rate'])
    assert r['vat_basis'] in ('incl.', 'excl.') and r['price_basis'] in ('per piece', 'per pack') and r['match'] in ('EXACT', 'CLOSE', 'APPROX'), r
order = {p: i for i, p in enumerate(NEW_NAMES)}; morder = {m: i for i, m in enumerate(MARKETS)}
rows.sort(key=lambda r: (order[r['product']], morder[r['market']], r['shop']))

wb = load_workbook(SRC)
# ---- widen ranges (idempotent) ----
reOff = re.compile(r'(\$[A-Z]{1,2}\$)(1000|3000)\b'); reRev = re.compile(r'(\$[A-Z]{1,2}\$)(216|600)\b')
nrep = 0
for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('='):
                v = reOff.sub(r'\g<1>%d' % OFF_MAX, c.value) if ('Offers!' in c.value or s.title == 'Offers') else c.value
                if 'Review!' in v: v = reRev.sub(r'\g<1>%d' % REV_MAX, v)
                if v != c.value: c.value = v; nrep += 1
print('formulas with widened ranges:', nrep)

def copy_style(src, dst):
    dst._style = copy.copy(src._style)
def set_table(ws, name, ref): ws.tables[name].ref = ref
def remap_cf(ws, fn):
    cf = ws.conditional_formatting; items = [(str(r.sqref), cf._cf_rules[r]) for r in list(cf._cf_rules)]
    cf._cf_rules.clear()
    for sq, rules in items:
        for rule in rules: ws.conditional_formatting.add(fn(sq), rule)
def last_row(ws, col=2):
    r = ws.max_row
    while r > 1 and ws.cell(r, col).value in (None, ''): r -= 1
    return r

# ---- Offers ----
ws = wb['Offers']; tmpl = 865
first = last_row(ws) + 1
VATTXT = {'Austria': 'AT 20%', 'Germany': 'DE 19%', 'Czech': 'CZ 21%', 'Slovakia': 'SK 23%', 'Romania': 'RO 21%'}
r = first
for o in rows:
    for col in range(2, 38): copy_style(ws.cell(tmpl, col), ws.cell(r, col))
    vat_txt = f"Price shown {o['vat_basis']} VAT ({VATTXT[o['market']]})" + ('; VAT removed for the net price' if o['vat_basis'] == 'incl.' else '')
    pack_txt = f"; pack of {o['pack_size']}, converted to per piece" if o['price_basis'] == 'per pack' else ''
    fetch_txt = '; page read with a headless browser (prices rendered by JavaScript)' if o['fetch_method'].startswith('browser') else ''
    vals = {2: o['product'], 3: o['market'], 4: o['shop'], 5: o['product_name'], 6: o['price_shown'], 7: o['currency'],
            8: f'=IF(AND(AE{r}="yes",ISNUMBER(AD{r})),AD{r},"")',
            9: f'=IF(AND(AE{r}="yes",ISNUMBER(AD{r}),ISNUMBER(INDEX(Sources!$G$7:$G$10,MATCH(AB{r},Sources!$B$7:$B$10,0)))),ROUND(AD{r}/INDEX(Sources!$G$7:$G$10,MATCH(AB{r},Sources!$B$7:$B$10,0)),6),"")',
            10: o['match'], 11: DATE_ISO, 13: f"{o['spec_txt']}; {o['matching_note']} {vat_txt}{pack_txt}{fetch_txt}. YES Gastro SKU {o['yes_sku']}.",
            14: o['url'], 15: o['brand'] or None, 16: o['vat_basis'], 17: o['vat_rate'], 18: o['pack_size'], 19: o['price_basis'], 20: o['match'],
            21: f'Added 2 Oct 2026 from live page (Claude research, {FAMILY_SHORT} expansion). ' + ('Net price read on page.' if o['vat_basis'] == 'excl.' else 'Gross price on page; VAT removed.'),
            24: 'Claude research 02.10.2026 (live shop pages)', 25: 'Claude', 26: o['stock'] or None, 27: 'Direct', 28: o['currency'], 29: 'yes',
            30: f'=IF(OR(F{r}="",Q{r}="",R{r}=""),"",ROUND(IF(P{r}="incl.",F{r}/(1+Q{r}),F{r})/IF(S{r}="per pack",R{r},1),2))',
            31: f'=IF(AND(AC{r}="yes",ISNUMBER(F{r}),IF(Compare!$H$5="All mapped",TRUE(),IF(Compare!$H$5="EXACT only",J{r}="EXACT",OR(J{r}="EXACT",J{r}="CLOSE")))),"yes","no")',
            32: f'=IF(ISNUMBER(H{r}),COUNTIFS($B$7:$B${OFF_MAX},B{r},$C$7:$C${OFF_MAX},C{r},$H$7:$H${OFF_MAX},"<"&H{r},$AH$7:$AH${OFF_MAX},"yes")+COUNTIFS($B$7:B{r},B{r},$C$7:C{r},C{r},$H$7:H{r},H{r},$AH$7:AH{r},"yes"),"")',
            33: f'=IF(ISNUMBER(I{r}),COUNTIFS($B$7:$B${OFF_MAX},B{r},$C$7:$C${OFF_MAX},C{r},$I$7:$I${OFF_MAX},"<"&I{r},$AI$7:$AI${OFF_MAX},"yes")+COUNTIFS($B$7:B{r},B{r},$C$7:C{r},C{r},$I$7:I{r},I{r},$AI$7:AI{r},"yes"),"")',
            34: f'=IF(ISNUMBER(H{r}),"yes","no")', 35: f'=IF(ISNUMBER(I{r}),"yes","no")', 36: DT,
            22: round((o['price_shown'] / (1 + o['vat_rate']) if o['vat_basis'] == 'incl.' else o['price_shown']) / (o['pack_size'] if o['price_basis'] == 'per pack' else 1), 2)}
    for col, v in vals.items(): ws.cell(r, col).value = v
    r += 1
last_off = r - 1
assert last_off < OFF_MAX
set_table(ws, 'BenchmarkOffers', f'B6:AK{last_off}')
remap_cf(ws, lambda sq: re.sub(r'(J7:J|K7:K)\d+', r'\g<1>%d' % OFF_MAX, sq))
print('Offers rows added:', last_off - first + 1, 'rows', first, '-', last_off)

# ---- Benchmarks ----
ws = wb['Benchmarks']; b0 = last_row(ws) + 1; br = b0
for p, c, sp, req in NEW:
    for col in range(2, 5): copy_style(ws.cell(b0 - 1, col), ws.cell(br, col))
    ws.cell(br, 2).value = p; ws.cell(br, 3).value = sp; ws.cell(br, 4).value = req; br += 1
set_table(ws, 'BenchmarkSpecifications', f'B6:D{br-1}')

# ---- Review ----
ws = wb['Review']; r0 = last_row(ws) + 1; rv = r0
def rev_formulas(r):
    B, C = f'$B{r}', f'$C{r}'
    cnt = lambda col: f'COUNTIFS(Offers!$B$7:$B${OFF_MAX},{B},Offers!$C$7:$C${OFF_MAX},{C},Offers!${col}$7:${col}${OFF_MAX},"yes")'
    pick = lambda val, rank, k: f'SUMIFS(Offers!${val}$7:${val}${OFF_MAX},Offers!$B$7:$B${OFF_MAX},{B},Offers!$C$7:$C${OFF_MAX},{C},Offers!${rank}$7:${rank}${OFF_MAX},{k})'
    f = {}
    for val, rank, flag, cols in (('H', 'AF', 'AH', (6, 7, 8)), ('I', 'AG', 'AI', (9, 10, 11))):
        n = cnt(flag)
        f[cols[0]] = f'=IF({n}=0,"",{pick(val, rank, 1)})'
        f[cols[1]] = f'=IF({n}=0,"",({pick(val, rank, f"ROUNDUP({n}/2,0)")}+{pick(val, rank, f"INT({n}/2)+1")})/2)'
        f[cols[2]] = f'=IF({n}=0,"",{pick(val, rank, n)})'
    f[15] = f'=IF(Compare!$E$5="Local prices",IF(F{r}="","",IF(Compare!$C$5="Minimum",F{r},IF(Compare!$C$5="Maximum",H{r},G{r}))),IF(I{r}="","",IF(Compare!$C$5="Minimum",I{r},IF(Compare!$C$5="Maximum",K{r},J{r}))))'
    return f
chk = rev_formulas(7)
for col in (6, 7, 8, 9, 10, 11, 15): assert ws.cell(7, col).value == chk[col], (col, ws.cell(7, col).value, chk[col])
for p, c, sp, req in NEW:
    for m in MARKETS:
        for col in range(2, 16): copy_style(ws.cell(r0 - 1, col), ws.cell(rv, col))
        ws.cell(rv, 2).value = p; ws.cell(rv, 3).value = m; ws.cell(rv, 4).value = CUR[m]; ws.cell(rv, 5).value = 'Net per piece'
        for col, fml in rev_formulas(rv).items(): ws.cell(rv, col).value = fml
        rv += 1
last_rev = rv - 1; assert last_rev < REV_MAX
set_table(ws, 'CountryReview', f'B6:O{last_rev}')
print('Review rows:', r0, '-', last_rev)

# ---- Compare ----
ws = wb['Compare']
L = next(rr for rr in range(12, ws.max_row + 1) if ws.cell(rr, 2).value == 'Research dates')
old_last = L - 2  # last product row
block = [[(copy.copy(ws.cell(rr, cc)._style), ws.cell(rr, cc).value) for cc in range(2, 9)] for rr in range(L - 1, L + 11)]
moved = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges if m.min_row >= L - 1]
for m in moved: ws.unmerge_cells(start_row=m[0], start_column=m[1], end_row=m[2], end_column=m[3])
for rr in range(L - 1, L + 12):
    for cc in range(2, 9): ws.cell(rr, cc).value = None
def cmp_row(r):
    return {ci + 3: (f'=IF(IF($E$5="Local prices",COUNTIFS(Offers!$B$7:$B${OFF_MAX},$B{r},Offers!$C$7:$C${OFF_MAX},{col}$10,Offers!$AH$7:$AH${OFF_MAX},"yes"),'
                     f'COUNTIFS(Offers!$B$7:$B${OFF_MAX},$B{r},Offers!$C$7:$C${OFF_MAX},{col}$10,Offers!$AI$7:$AI${OFF_MAX},"yes"))=0,"—",'
                     f'SUMIFS(Review!$O$7:$O${REV_MAX},Review!$B$7:$B${REV_MAX},$B{r},Review!$C$7:$C${REV_MAX},{col}$10))') for ci, col in enumerate('CDEFGH')}
chk = cmp_row(12)
for col in range(3, 9): assert ws.cell(12, col).value == chk[col], (col, ws.cell(12, col).value, chk[col])
cr = L - 1
for p, c, sp, req in NEW:
    for col in range(2, 9): copy_style(ws.cell(old_last, col), ws.cell(cr, col))
    ws.cell(cr, 2).value = p
    for col, fml in cmp_row(cr).items(): ws.cell(cr, col).value = fml
    cr += 1
last_cmp = cr - 1
newL = last_cmp + 2; shift = newL - L
for i, rowvals in enumerate(block):
    rr = L - 1 + i + shift
    for j, (style, val) in enumerate(rowvals):
        cell = ws.cell(rr, j + 2); cell._style = style
        if isinstance(val, str) and val.startswith('='):
            val = re.sub(r'\bB(\d+)\b', lambda m: f'B{int(m.group(1)) + shift}' if L + 2 <= int(m.group(1)) <= L + 7 else m.group(0), val)
        cell.value = val
for m in moved: ws.merge_cells(start_row=m[0] + shift, start_column=m[1], end_row=m[2] + shift, end_column=m[3])
remap_cf(ws, lambda sq: f'C12:H{last_cmp}' if sq.startswith('C12:H') else (f'D{L+2+shift}:D{L+7+shift}' if sq == f'D{L+2}:D{L+7}' else sq))
ws['B3'].value = str(ws['B3'].value).replace(' Net prices per piece.', '') + f' {spec["compare_subtitle"]} Net prices per piece.'
print('Compare product rows', L - 1, '-', last_cmp, '; dates block from', newL)

# ---- Pricing ----
ws = wb['Pricing']
P = next(rr for rr in range(15, ws.max_row + 1) if ws.cell(rr, 2).value == 'How to read it')
old_last = P - 2
notes = [[(copy.copy(ws.cell(rr, cc)._style), ws.cell(rr, cc).value) for cc in range(2, 31)] for rr in range(P - 1, P + 8)]
moved = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges if m.min_row >= P - 1]
for m in moved: ws.unmerge_cells(start_row=m[0], start_column=m[1], end_row=m[2], end_column=m[3])
tmpl_f = {cc: ws.cell(old_last, cc).value for cc in range(2, 68)}; tmpl_s = {cc: copy.copy(ws.cell(old_last, cc)._style) for cc in range(2, 68)}
for rr in range(P - 1, P + 9):
    for cc in range(2, 68): ws.cell(rr, cc).value = None
reref = re.compile(r'(\$?)([A-Z]{1,2})(\$?)%d(?![0-9])' % old_last)
pr = P - 1
for p, c, sp, req in NEW:
    for cc in range(2, 68):
        cell = ws.cell(pr, cc); cell._style = copy.copy(tmpl_s[cc]); v = tmpl_f[cc]
        if isinstance(v, str) and v.startswith('='): v = reref.sub(lambda m: f'{m.group(1)}{m.group(2)}{m.group(3)}{pr}', v)
        elif cc == 2: v = p
        elif cc == 3: v = c
        elif cc == 4: v = None
        cell.value = v
    pr += 1
last_pr = pr - 1; nb = last_pr + 2; shift = nb - P
for i, rowvals in enumerate(notes):
    rr = P - 1 + i + shift
    for j, (style, val) in enumerate(rowvals):
        cell = ws.cell(rr, j + 2); cell._style = style; cell.value = val
for m in moved: ws.merge_cells(start_row=m[0] + shift, start_column=m[1], end_row=m[2] + shift, end_column=m[3])
ws.cell(nb + 7, 2).value = str(ws.cell(nb + 7, 2).value) + f' {FAMILY} expansion of 2 Oct 2026 (rows {P-1}-{last_pr}): cost empty for the same reason.'
remap_cf(ws, lambda sq: re.sub(r'^([A-Z]{1,2})15:([A-Z]{1,2})%d$' % old_last, lambda m: f'{m.group(1)}15:{m.group(2)}{last_pr}', sq))
print('Pricing rows', P - 1, '-', last_pr, '; notes from', nb)

# ---- Change log ----
ws = wb['Change log']; cl = last_row(ws) + 1; cl0 = cl
def log(*vals):
    global cl
    for col in range(2, 13): copy_style(ws.cell(cl0 - 1, col), ws.cell(cl, col))
    for col, v in zip(range(2, 13), (TODAY,) + vals): ws.cell(cl, col).value = v
    cl += 1
log('Offers / Review / Compare / Pricing', 'all formulas', 'All', '—', 'All benchmarks', 'Range limits in formulas and conditional formats',
    'Offers rows 7-3000; Review rows 7-600', f'Offers rows 7-{OFF_MAX}; Review rows 7-{REV_MAX}', '—', f'Widened again for the {FAMILY} expansion; values unchanged.')
log('Benchmarks', f'{b0}-{br-1}', 'AT / DE / CZ / SK / RO', '—', f'{len(NEW)} {FAMILY} variants', 'Rows added', '—', f'{len(NEW)} benchmark specifications', 'Shopify "Sauce Dispensers" collection (read-only), 2 Oct 2026', spec['benchmark_reason'])
log('Review', f'{r0}-{last_rev}', 'All', '—', f'{len(NEW)} {FAMILY} variants', 'Rows added', '—', f'{last_rev-r0+1} product x market rows', '—', 'Standard min/median/max formulas.')
log('Compare', f'{L-1}-{last_cmp}', 'All', '—', f'{len(NEW)} {FAMILY} variants', 'Rows added; research-dates block moved', f'Research dates at rows {L}-{L+10}', f'Research dates at rows {newL}-{newL+10}', '—', 'Standard formulas; the research-date block and note were moved below the new rows.')
log('Pricing', f'{P-1}-{last_pr}', 'All', '—', f'{len(NEW)} {FAMILY} variants', 'Rows added; notes moved', f'Notes at rows {P}-{P+7}', f'Notes at rows {nb}-{nb+7}', '—', f'Market-entry logic copied from row {old_last}. Cost (column D) left empty: the Gastroplast price list was not supplied.')
cnt = defaultdict(Counter)
for o in rows: cnt[o['product']][o['market']] += 1
for p, c, sp, req in NEW:
    parts = [f'{m} {cnt[p][m]}' for m in ('Austria', 'Germany', 'Czech', 'Slovakia', 'Romania') if cnt[p][m]]; tot = sum(cnt[p].values())
    log('Offers', 'appended', 'AT / DE / CZ / SK / RO', 'see rows', p, 'Offers added', '—', f'{tot} offers ({", ".join(parts) if parts else "none found"})', 'Live shop pages 2 Oct 2026 (URLs in Offers)', f'{FAMILY} expansion; YES Gastro code {c}.' + (' No usable offer found in any market.' if tot == 0 else ''))
if EXTRA:
    for e in csv.DictReader(open(EXTRA, encoding='utf-8')): log(e['sheet'], e['row'], e['market'], e['competitor'], e['product'], e['field'], e['old'], e['new'], e['evidence'], e['reason'])
set_table(ws, 'ChangeLog', f'B5:L{cl-1}')
ws['B3'].value = str(ws['B3'].value).replace(' Only listed fields changed', f' and the 2 Oct 2026 {FAMILY} expansion. Only listed fields changed') if FAMILY not in str(ws['B3'].value) else ws['B3'].value
ws2 = wb['Sources']; ws2['C18'].value = str(ws2['C18'].value).replace(' Other markets show a dash', f' The {len(NEW)} {FAMILY} variants added on 2 Oct 2026 have the same five-market coverage. Other markets show a dash')
print('Change log rows', cl0, '-', cl - 1)
wb.save(OUT); print('saved', OUT)
