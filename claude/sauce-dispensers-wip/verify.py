import sys, math, statistics
from decimal import Decimal, ROUND_HALF_UP
def sig15(x): return float(f'{x:.15g}')
def xround(x, n): return float(Decimal(f'{x:.15g}').quantize(Decimal(1).scaleb(-n), rounding=ROUND_HALF_UP))
from openpyxl import load_workbook
from collections import defaultdict
F = sys.argv[1]
wbv = load_workbook(F, data_only=True); wbf = load_workbook(F)
# 0. formula errors + formula text left as values
errs = []; ftext = []
for ws in wbv.worksheets:
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str):
                if v.startswith('#') and v.rstrip('!?/0ADEFILMNRUV!').upper() in ('#','') or v in ('#VALUE!','#REF!','#NAME?','#DIV/0!','#N/A','#NUM!','#NULL!'): errs.append((ws.title, c.coordinate, v))
                if v.startswith('='): ftext.append((ws.title, c.coordinate, v[:60]))
print('formula errors:', len(errs), errs[:10]); print('formula text as values:', len(ftext), ftext[:10])
# 1. FX
src = wbv['Sources']; rate = {src.cell(r, 2).value: src.cell(r, 7).value for r in range(7, 11)}
print('units per EUR:', rate)
cmp = wbv['Compare']; sel, view, filt = cmp['C5'].value, cmp['E5'].value, cmp['H5'].value
print('Compare selectors:', sel, view, filt)
# 2. recompute offers net/EUR for new rows
off = wbv['Offers']; offf = wbf['Offers']
new_products = set()
bm = wbv['Benchmarks']
for r in range(42, bm.max_row + 1):
    if bm.cell(r, 2).value: new_products.add(bm.cell(r, 2).value)
groups = defaultdict(list); mism = []
nrows = 0
for r in range(7, off.max_row + 1):
    p, m = off.cell(r, 2).value, off.cell(r, 3).value
    if p not in new_products: continue
    nrows += 1
    F_, P, Q, R, S, J, AB, AC = (off.cell(r, c).value for c in (6, 16, 17, 18, 19, 10, 28, 29))
    net = xround((F_ / (1 + Q) if P == 'incl.' else F_) / (R if S == 'per pack' else 1), 2)
    eur = xround(net / rate[AB], 6)
    inc = AC == 'yes' and (filt == 'All mapped' or (filt == 'EXACT only' and J == 'EXACT') or (filt == 'EXACT + CLOSE' and J in ('EXACT', 'CLOSE')))
    H, I, AD, AE = (off.cell(r, c).value for c in (8, 9, 30, 31))
    if abs(AD - net) > 1e-9: mism.append(('Offers AD', r, AD, net))
    if inc:
        if not (isinstance(H, (int, float)) and abs(H - net) < 1e-9): mism.append(('Offers H', r, H, net))
        if not (isinstance(I, (int, float)) and abs(I - eur) < 1e-9): mism.append(('Offers I', r, I, eur))
        groups[(p, m)].append((net, eur))
    else:
        if H not in (None, ''): mism.append(('Offers H should be blank', r, H))
print('new offer rows checked:', nrows)
# 3. Review
rev = wbv['Review']; revstats = {}
for r in range(217, rev.max_row + 1):
    p, m = rev.cell(r, 2).value, rev.cell(r, 3).value
    if not p: continue
    vals = groups.get((p, m), [])
    got = [rev.cell(r, c).value for c in (6, 7, 8, 9, 10, 11)]
    if not vals:
        exp = [None] * 6
        if any(g not in (None, '') for g in got): mism.append(('Review nonblank', r, got))
        continue
    loc = sorted(v[0] for v in vals); eu = sorted(v[1] for v in vals)
    exp = [loc[0], statistics.median(loc), loc[-1], eu[0], statistics.median(eu), eu[-1]]
    revstats[(p, m)] = exp
    for g, e, lab in zip(got, exp, 'FGHIJK'):
        if not isinstance(g, (int, float)) or abs(g - e) > 1e-6: mism.append(('Review ' + lab, r, p, m, g, e))
# 4. Compare (view = Net per piece EUR, statistic = sel)
for r in range(47, cmp.max_row + 1):
    p = cmp.cell(r, 2).value
    if p not in new_products: continue
    for ci, m in zip(range(3, 9), ['Hungary', 'Romania', 'Austria', 'Germany', 'Czech', 'Slovakia']):
        got = cmp.cell(r, ci).value
        if (p, m) not in revstats:
            if got != '—': mism.append(('Compare', r, m, got, '—'))
            continue
        st = revstats[(p, m)]
        base = st[3:] if view == 'Net per piece EUR' else st[:3]
        e = base[0] if sel == 'Minimum' else base[2] if sel == 'Maximum' else base[1]
        if not isinstance(got, (int, float)) or abs(got - e) > 1e-6: mism.append(('Compare', r, m, got, e))
# 5. Pricing
pr = wbv['Pricing']
C5, C6, C7, C8, C9, D9, E9, C10 = (pr[c].value for c in ('C5', 'C6', 'C7', 'C8', 'C9', 'D9', 'E9', 'C10'))
blocks = {'Austria': 5, 'Germany': 15, 'Czech': 25, 'Slovakia': 36, 'Romania': 46, 'Hungary': 57}
def floor_step(x, step): return math.floor(x / step + 1e-9) * step
flags = []
for r in range(50, pr.max_row + 1):
    p = pr.cell(r, 2).value
    if p not in new_products: continue
    cost = pr.cell(r, 4).value
    for m, c0 in blocks.items():
        n = len(groups.get((p, m), []))
        gotE = pr.cell(r, c0).value
        if gotE != n: mism.append(('Pricing count', r, m, gotE, n))
        extra = 1 if m in ('Czech', 'Romania', 'Hungary') else 0  # extra local-currency column after Entry price EUR
        if n == 0:
            if pr.cell(r, c0 + 5).value != '—': mism.append(('Pricing entry blank', r, m, pr.cell(r, c0 + 5).value))
            continue
        st = revstats[(p, m)]; cheap, med, mx = st[3], st[4], st[5]
        ex = C8 if n < 3 else 0
        if sig15(cheap) < sig15(med * (1 - C7)): raw = med * (1 - C5 - ex); basis = 'median (cheapest is an outlier)'
        else:
            raw = min(med * (1 - C5 - ex), cheap * (1 - C6 - ex)); basis = 'undercut cheapest' if sig15(cheap * (1 - C6)) < sig15(med * (1 - C5)) else 'median'
        step = C9 if raw < 10 else D9 if raw < 100 else E9
        entry = floor_step(raw, step)
        if isinstance(cost, (int, float)) and cost * (1 + C10) > entry: entry = round(cost * (1 + C10), 2)
        gotB, gotJ = pr.cell(r, c0 + 4).value, pr.cell(r, c0 + 5).value
        if gotB != basis: mism.append(('Pricing basis', r, m, gotB, basis))
        if not isinstance(gotJ, (int, float)) or abs(gotJ - entry) > 1e-6: mism.append(('Pricing entry', r, m, gotJ, entry))
        for lab, off_, e in (('cheapest', 1, cheap), ('median', 2, med), ('max', 3, mx)):
            g = pr.cell(r, c0 + off_).value
            if not isinstance(g, (int, float)) or abs(g - e) > 1e-6: mism.append(('Pricing ' + lab, r, m, g, e))
        if isinstance(cost, (int, float)):
            mk = entry / cost - 1
            if mk < 0.5: flags.append((p, m, entry, cost, round(mk, 3)))
print('MISMATCHES:', len(mism))
for x in mism[:40]: print('  ', x)
print('thin-margin flags (<+50%):', flags)
